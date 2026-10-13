"""
Label the agent's decisions to find what to fix in the agent.

    python -m evaluation.make_agent_worksheet \
        --agent test_reports/chatgpt_full_Qwen2.5-32B-Instruct-AWQ.json \
        --answers datasets/candidates_wildchat.json datasets/candidates_hc3.json \
        [--rules test_reports/chatgpt_rules_only.json] \
        --out test_reports/agent_worksheet.xlsx

Row order (label top-down; you can stop after group 2):
  1  agent said met            -> is it a false positive?
  2  agent said not_met / didn't run, but the rules flagged it -> is it a miss?
  3  agent undetermined
  4  everything else the agent said not_met (spot-check for misses)
"""

import argparse
import json

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

CRIT = {"C1": "Statistic with no specific source nearby?",
        "C2": "Does a reliable source contradict the number?",
        "C3": "Slanted / loaded wording where neutral would do?",
        "C4": "Sweeping absolute claim the evidence can't support?",
        "C5": "'Experts say' instead of a named source?",
        "C6": "Cause-and-effect stated without evidence?"}
GROUP = {1: "1 agent met: false positive?", 2: "2 possible miss (rules flagged)",
         3: "3 undetermined", 4: "4 agent not_met: spot-check"}
F = "Arial"


def load(path):
    out = {}
    if path:
        for r in json.load(open(path, encoding="utf-8"))["answers_detail"]:
            for c in r["claims"]:
                for k, v in c["criteria"].items():
                    out[(r["id"], c["sentence"], k)] = (r.get("domain", ""), v)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--agent", required=True)
    ap.add_argument("--rules")
    ap.add_argument("--answers", nargs="*", default=[])
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    agent, rules = load(a.agent), load(a.rules)
    q = {it["id"]: it.get("question", "") for p in a.answers for it in json.load(open(p, encoding="utf-8"))}

    rows = []
    for key, (dom, v) in agent.items():
        rflag = rules.get(key, ("", ""))[1] == "met"
        g = 1 if v == "met" else 3 if v == "undetermined" else 2 if rflag else 4
        rows.append((g, key[0], key[2], dom, key[1], v))
    for key, (dom, v) in rules.items():  # rules flagged, agent never checked it
        if v == "met" and key not in agent:
            rows.append((2, key[0], key[2], dom, key[1], "not run"))
    rows.sort(key=lambda r: (r[0], r[2], r[1]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Label"
    ws.append(["#", "Group", "Answer ID", "Domain", "Question", "Sentence", "Criterion",
               "Criterion question", "Agent decision", "Human label", "Notes (why is the agent wrong?)",
               "Agent correct?", "Error type"])
    for i, (g, aid, crit, dom, sent, v) in enumerate(rows, start=2):
        ws.append([i - 1, GROUP[g], aid, dom, q.get(aid, ""), sent, crit, CRIT.get(crit, ""), v, None, None,
                   f'=IF(OR(J{i}="",J{i}="unsure"),"",IF(I{i}=J{i},1,IF(AND(I{i}="not run",J{i}="not_met"),1,0)))',
                   f'=IF(OR(J{i}="",J{i}="unsure"),"",IF(AND(I{i}="met",J{i}="not_met"),"false positive",'
                   f'IF(AND(I{i}<>"met",J{i}="met"),"miss","")))'])
    n = len(rows) + 1

    thin, hfill, yellow = Side(style="thin", color="D9D9D9"), PatternFill("solid", fgColor="1F3864"), PatternFill("solid", fgColor="FFF2CC")
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(name=F, bold=True, color="FFFFFF"), hfill, Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2, max_row=n):
        for c in row:
            c.font, c.alignment, c.border = Font(name=F, size=10), Alignment(wrap_text=True, vertical="top"), Border(bottom=thin)
        row[9].fill = row[10].fill = yellow
    for col, w in zip("ABCDEFGHIJKLM", [5, 18, 10, 11, 28, 60, 7, 26, 11, 12, 32, 9, 13]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = f"A1:M{n}"
    dv = DataValidation(type="list", formula1='"met,not_met,unsure"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"J2:J{n}")
    ws.conditional_formatting.add(f"M2:M{n}", FormulaRule(formula=['M2<>""'], fill=PatternFill("solid", fgColor="F8CBAD")))

    s = wb.create_sheet("Summary")
    s.append(["Criterion", "Question", "Agent says met", "Labeled", "Agent accuracy",
              "False positives", "Misses", "Precision (of agent met)"])
    G, Lc, Lm = f"Label!$G$2:$G${n}", f"Label!$L$2:$L${n}", f"Label!$M$2:$M${n}"
    I = f"Label!$I$2:$I${n}"
    J = f"Label!$J$2:$J${n}"
    for j, c in enumerate(sorted(CRIT), start=2):
        s.append([c, CRIT[c], f'=COUNTIFS({G},A{j},{I},"met")',
                  f'=COUNTIFS({G},A{j},{Lc},1)+COUNTIFS({G},A{j},{Lc},0)',
                  f'=IFERROR(COUNTIFS({G},A{j},{Lc},1)/D{j},"")',
                  f'=COUNTIFS({G},A{j},{Lm},"false positive")',
                  f'=COUNTIFS({G},A{j},{Lm},"miss")',
                  f'=IFERROR(COUNTIFS({G},A{j},{I},"met",{J},"met")/(COUNTIFS({G},A{j},{I},"met",{J},"met")+F{j}),"")'])
    t = len(CRIT) + 2
    s.append(["All", "", f"=SUM(C2:C{t-1})", f"=SUM(D2:D{t-1})",
              f'=IFERROR(SUMPRODUCT(E2:E{t-1},D2:D{t-1})/D{t},"")', f"=SUM(F2:F{t-1})", f"=SUM(G2:G{t-1})",
              f'=IFERROR(COUNTIFS({I},"met",{J},"met")/(COUNTIFS({I},"met",{J},"met")+F{t}),"")'])
    for c in s[1]:
        c.font, c.fill, c.alignment = Font(name=F, bold=True, color="FFFFFF"), hfill, Alignment(wrap_text=True)
    for row in s.iter_rows(min_row=2, max_row=t):
        for c in row:
            c.font = Font(name=F, bold=(row[0].value == "All"))
        row[4].number_format = row[7].number_format = "0.0%"
    for col, w in zip("ABCDEFGH", [10, 46, 12, 10, 13, 13, 10, 14]):
        s.column_dimensions[col].width = w
    counts = {g: sum(r[0] == g for r in rows) for g in GROUP}
    notes = ["How to use",
             "Fill the yellow 'Human label' on the Label sheet: met = the problem IS in the sentence, not_met = it isn't, unsure = skip.",
             f"Work top-down. Group 1 ({counts[1]} rows): every 'met' the agent gave — mark false positives. "
             f"Group 2 ({counts[2]} rows): the rules flagged it but the agent didn't — mark misses.",
             f"Groups 3-4 ({counts[3]}+{counts[4]} rows) are optional spot-checks.",
             "In 'Notes', say WHY the agent is wrong (e.g. 'advice, not a generalization'); these notes become the prompt fixes.",
             f"Source: {a.agent}"]
    for k, line in enumerate(notes):
        s.cell(t + 2 + k, 1, line).font = Font(name=F, bold=(k == 0), italic=(k == len(notes) - 1), color="404040")
    wb.move_sheet("Summary", offset=-1)
    wb.save(a.out)
    print(f"{len(rows)} rows; groups: {counts} -> {a.out}")


if __name__ == "__main__":
    main()
