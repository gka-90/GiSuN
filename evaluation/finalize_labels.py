"""
Clean combined label sheet after the Gideon/Suning adjudication meeting (2026-10-27): one final label
per row, how it was decided, the agent's correctness, and a summary computed by formulas.

    python -m evaluation.finalize_labels --csv "<resolved combined CSV>" --out test_reports/GiSuN_labels_final.xlsx

Input: the combined sheet exported as CSV, with "Combined label" filled for every row and the meeting
notes in "Suning and Gideon notes". Reviewers see the sentence and the 2 before it, not the full answer.
"""

import argparse
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

F = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
CRITERIA = ["C1", "C3", "C4", "C5", "C6"]
FP_TYPES = ["FP-COMMON", "FP-ATTRIBUTED", "FP-NEUTRAL", "FP-OTHER", "FP-SOURCE", "FP-NOT-STAT", "FP-CITATION",
            "MISS-JUDGED", "(not given)"]
# Rows the meeting settled as not_met because the reader's prior knowledge is enough, with no category recorded
PRIOR_KNOWLEDGE_ROWS = {"25", "26", "27", "28"}


def _cell(ws, r, c, v, fill=None, bold=False, color=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font, cell.alignment, cell.border = Font(name=F, bold=bold, color=color), WRAP, THIN
    if fill:
        cell.fill = fill
    return cell


def _header(ws, cols, widths, row=1):
    for c, (h, w) in enumerate(zip(cols, widths), 1):
        _cell(ws, row, c, h, HEAD, bold=True, color="FFFFFF")
        ws.column_dimensions[ws.cell(row=row, column=c).column_letter].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def how_decided(r):
    if r["Labeled by"] != "Both":
        return f"single labeler ({r['Labeled by']})"
    if r["Agreement"] == "agree":
        return "agreed when labeling"
    side = "Gideon" if r["Combined label"] == r["Gideon label"] else "Suning"
    return f"settled in meeting (sided with {side})"


def error_type(r):
    """The false-positive type for rows where the agent said met and the final label is not_met."""
    if not (r["Agent decision"] == "met" and r["Combined label"] == "not_met"):
        return ""
    t = r["Error category (combined)"] or r["Gideon error category"] or r["Suning error category"]
    if not t and r["#"] in PRIOR_KNOWLEDGE_ROWS:
        t = "FP-COMMON"  # meeting note: "prior knowledge, that's why it's not met"
    return t or "(not given)"


def build(rows, out):
    wb = Workbook()
    lab = wb.active
    lab.title = "Labels"
    cols = ["#", "Answer ID", "Domain", "Question", "Previous 2 sentences", "Sentence", "Criterion", "Agent decision",
            "Final label", "How decided", "Agent correct?", "Error type", "Notes"]
    _header(lab, cols, [5, 9, 13, 30, 40, 55, 9, 12, 10, 26, 9, 15, 40])
    for i, r in enumerate(rows, 2):
        notes = " | ".join(x for x in (r["Notes (both labelers)"], r.get("Suning and Gideon notes", "")) if x.strip())
        vals = [int(r["#"]), r["Answer ID"], r["Domain"], r["Question"], r["Source of sentences"], r["Sentence"],
                r["Criterion"], r["Agent decision"], r["Combined label"], how_decided(r),
                f'=IF(H{i}="met",IF(I{i}="met",1,0),IF(I{i}="met",0,1))', error_type(r), notes]
        for c, v in enumerate(vals, 1):
            _cell(lab, i, c, v)
    last = len(rows) + 1
    L = lambda col: f"Labels!${col}$2:${col}${last}"

    s = wb.create_sheet("Summary", 0)
    s.column_dimensions["A"].width = 30
    for col in "BCDEFGHIJ":
        s.column_dimensions[col].width = 13
    _cell(s, 1, 1, "Agent vs final labels (184 rows, after the 2026-10-27 meeting)", bold=True)
    hdr = ["Criterion", "Rows", "Agent said met", "Correct", "Precision", "Undetermined", "Undetermined that were met",
           "Final met"]
    for c, h in enumerate(hdr, 1):
        _cell(s, 2, c, h, HEAD, bold=True, color="FFFFFF")
    for k, cid in enumerate(CRITERIA + ["All"], 3):
        crit = f'{L("G")},"{cid}",' if cid != "All" else ""
        vals = [cid, f'=COUNTIFS({crit}{L("A")},">0")', f'=COUNTIFS({crit}{L("H")},"met")',
                f'=COUNTIFS({crit}{L("H")},"met",{L("I")},"met")', f'=IFERROR(D{k}/C{k},"")',
                f'=COUNTIFS({crit}{L("H")},"undetermined")', f'=COUNTIFS({crit}{L("H")},"undetermined",{L("I")},"met")',
                f'=COUNTIFS({crit}{L("I")},"met")']
        for c, v in enumerate(vals, 1):
            cell = _cell(s, k, c, v, bold=(cid == "All"))
            if c == 5:
                cell.number_format = "0.00"
    r0 = 3 + len(CRITERIA) + 2
    _cell(s, r0, 1, "False positives by error type (agent said met, final label not_met)", bold=True)
    for c, h in enumerate(["Error type"] + CRITERIA + ["All"], 1):
        _cell(s, r0 + 1, c, h, HEAD, bold=True, color="FFFFFF")
    for k, t in enumerate(FP_TYPES, r0 + 2):
        _cell(s, k, 1, t)
        for c, cid in enumerate(CRITERIA + ["All"], 2):
            crit = f'{L("G")},"{cid}",' if cid != "All" else ""
            _cell(s, k, c, f'=COUNTIFS({crit}{L("L")},"{t}")')
    r1 = r0 + 2 + len(FP_TYPES) + 1
    _cell(s, r1, 1, "How the final labels were decided", bold=True)
    for k, h in enumerate(["agreed when labeling", "settled in meeting (sided with Gideon)",
                           "settled in meeting (sided with Suning)", "single labeler (Gideon)",
                           "single labeler (Suning)"], r1 + 1):
        _cell(s, k, 1, h)
        _cell(s, k, 2, f'=COUNTIF({L("J")},"{h}")')

    m = wb.create_sheet("Meeting decisions")
    _header(m, ["#", "Criterion", "Sentence", "Gideon", "Suning", "Final", "Sided with", "Meeting note"],
            [5, 9, 60, 9, 9, 9, 11, 60])
    disputed = [r for r in rows if r["Agreement"] == "DISAGREE"]
    for i, r in enumerate(disputed, 2):
        side = "Gideon" if r["Combined label"] == r["Gideon label"] else "Suning"
        for c, v in enumerate([int(r["#"]), r["Criterion"], r["Sentence"], r["Gideon label"], r["Suning label"],
                               r["Combined label"], side, r.get("Suning and Gideon notes", "")], 1):
            _cell(m, i, c, v)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--csv", required=True)
    ap.add_argument("--out", default="test_reports/GiSuN_labels_final.xlsx")
    a = ap.parse_args()
    rows = list(csv.DictReader(open(a.csv, encoding="utf-8")))
    missing = [r["#"] for r in rows if r["Combined label"] not in ("met", "not_met")]
    if missing:
        raise SystemExit(f"rows without a final label: {', '.join(missing)}")
    build(rows, a.out)
    print(f"wrote {a.out}: {len(rows)} rows")
