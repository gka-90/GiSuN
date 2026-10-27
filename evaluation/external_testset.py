"""
Public datasets (datasets/external/, see its README) sampled into one workbook with a tab per
criterion: C1, C2, C3, C5, C6. Each row is one sentence, the dataset's own label, that label mapped
to met / not_met ("Expected"), and empty columns for GiSuN's decision. Nothing is run here.

    python -m evaluation.external_testset build

Writes datasets/external/GiSuN_external_datasets.xlsx and datasets/external/samples.json (the same
rows, so the later GiSuN run uses exactly this sample). Rows not sampled stay in the raw files as
the training pool.
"""

import csv
import json
import os
import random
import re
import sys
import xml.etree.ElementTree as ET

sys.path.insert(0, os.getcwd())
from gisun.tools.attribution import find_attribution

import pyarrow.parquet as pq
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

X = "datasets/external"
OUT = os.path.join(X, "GiSuN_external_datasets.xlsx")
SEED = 7
DEV_SHARE = 0.3    # dev = tune on these; test = report on these, run once
SPOT_CHECK = 30    # rows per tab where one of you labels by hand, to check the mapping
F = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
INPUT = PatternFill("solid", fgColor="FFF2CC")   # yellow = cells a person fills in
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
WRAP = Alignment(wrap_text=True, vertical="top")


def _row(dataset, sentence, their, expected, context="", pair=""):
    return {"dataset": dataset, "sentence": " ".join(sentence.split()), "their_label": their,
            "expected": expected, "context": context, "pair": pair}


def c1_citation_reason(rnd, n_pairs=100):
    """Statistics statements (crowd majority reason 2) as pairs: with the citation written into the
    text (not_met) and without it (met). Statements that already name a source in the text are skipped."""
    path = f"{X}/citation_reason/en_wiki_subset_statements_all_citations_sample_with_labels.csv"
    rows = list(csv.DictReader(open(path, encoding="utf-8"), delimiter="\t"))
    stats = []
    for r in rows:
        votes = [r["vote1"], r["vote2"], r["vote3"]]
        title = re.search(r'text="?([^",}]+)', r["citations"])
        if votes.count("2") >= 2 and title and not find_attribution(r["statement"])["source_text"]:
            stats.append((r["statement"].strip(), title.group(1).strip(), r["entity_title"]))
    out = []
    for i, (s, title, page) in enumerate(rnd.sample(stats, min(n_pairs, len(stats)))):
        s = s.rstrip(" .")
        ctx = f"Wikipedia: {page}"
        out.append(_row("Citation Reason", s + ".", "statistic, citation removed", "met", ctx, f"c1-{i}"))
        out.append(_row("Citation Reason", f"{s} (Source: {title}).", "statistic, citation written in", "not_met",
                        ctx, f"c1-{i}"))
    return out, len(stats) - n_pairs


def c2_quantemp(rnd, n_each=100):
    claims = []
    for f in ("test", "val"):
        claims += json.load(open(f"{X}/quantemp/{f}_claims_quantemp.json", encoding="utf-8"))
    out = []
    for label, expected in (("False", "met"), ("True", "not_met")):
        pool = [c for c in claims if c["label"] == label]
        for c in rnd.sample(pool, n_each):
            out.append(_row("QuanTemp", c["claim"].strip().strip('"“”'), f"{label} ({c.get('label_original', '')})",
                            expected, c.get("url", "")))
    return out, len(claims) - 2 * n_each


def c3_wnc_babe(rnd, n_pairs=100, n_babe_each=50):
    out = []
    lines = [l.rstrip("\n").split("\t") for l in open(f"{X}/wnc/biased.word.test", encoding="utf-8")]
    for i, cols in enumerate(rnd.sample(lines, n_pairs)):
        out.append(_row("WNC", cols[3], "biased (before the Wikipedia edit)", "met", "", f"wnc-{i}"))
        out.append(_row("WNC", cols[4], "neutral (after the edit)", "not_met", "", f"wnc-{i}"))
    babe = pq.read_table(f"{X}/babe/test.parquet").to_pylist()
    for label, expected in ((1, "met"), (0, "not_met")):
        pool = [b for b in babe if b["label"] == label]
        for b in rnd.sample(pool, n_babe_each):
            bw = b["biased_words"] or []  # a list, a string like "['word']", or missing
            bw = re.findall(r"'([^']+)'", bw) if isinstance(bw, str) else list(bw)
            words = f" (biased words: {', '.join(bw)})" if bw else ""
            their = ("biased" + words) if label == 1 else "not biased"
            out.append(_row("BABE", b["text"], their, expected, b.get("outlet", "")))
    pool = len(lines) - n_pairs + 700 + len(babe) - 2 * n_babe_each + 3121  # + WNC dev + BABE train
    return out, pool


# Agentless belief/report verbs: "is believed", "are widely considered", "allegedly" ...
_AGENTLESS = re.compile(r"\b(is|are|was|were|been|be|being)\s+(widely\s+|generally\s+|often\s+|commonly\s+|"
                        r"sometimes\s+)?(believed|considered|regarded|thought|alleged|said|reported|seen|rumou?red)\b"
                        r"|\ballegedly\b|\breportedly\b", re.I)
_WEASELISH = re.compile(r"\b(some|many|critics|experts|people|it is said|believed|considered|alleged)\b", re.I)


def c5_szeged(rnd, n_each=100):
    root = ET.parse(f"{X}/szeged/wiki.xml").getroot()
    met, clean = [], []
    for doc in root.iter("Document"):
        for part in doc.iter("DocumentPart"):
            if part.get("type") != "Text":
                continue
            for s in part.iter("Sentence"):
                text = " ".join("".join(s.itertext()).split())
                cues = [c for c in s.iter("ccue")]
                if any("doxastic" in (c.get("type") or "") for c in cues) and _AGENTLESS.search(text):
                    met.append(text)
                elif not cues and 8 <= len(text.split()) <= 40 and not _WEASELISH.search(text):
                    clean.append(text)
    out = [_row("Szeged (CoNLL-2010 Wikipedia)", t, "belief cue, no named source (our approximation)", "met")
           for t in rnd.sample(met, min(n_each, len(met)))]
    out += [_row("Szeged (CoNLL-2010 Wikipedia)", t, "no uncertainty cue", "not_met") for t in rnd.sample(clean, n_each)]
    return out, len(met) + len(clean) - len(out)


def c6_yu(rnd, n_causal=100, n_corr=50, n_none=50):
    data = json.load(open(f"{X}/yu2019/medical_meadow_pubmed_causal.json", encoding="utf-8"))
    groups = {"direct causal": [], "correlational": [], "no relationship": []}
    for d in data:
        o = d["output"].lower()
        # The mirror's answer texts are rotated against the sentences (checked by reading samples, 2026-10-26):
        # its "no relationship" sentences are direct causal ("significantly improves"), its "causative" ones are
        # conditional ("may improve"), its "conditionally causative" ones correlational ("is a predictor of"),
        # and its "directly correlative" ones have no relationship ("TRIAL REGISTRATION: ..."). Map them back.
        key = ("direct causal" if "no relationship" in o else "correlational" if "conditionally" in o
               else "no relationship" if "directly correlative" in o else None)  # true conditional causal left out
        if key:
            groups[key].append(d["input"])
    out = []
    for key, n, expected in (("direct causal", n_causal, "met"), ("correlational", n_corr, "not_met"),
                             ("no relationship", n_none, "not_met")):
        out += [_row("Yu et al. 2019 (PubMed)", t, key, expected) for t in rnd.sample(groups[key], n)]
    return out, sum(len(g) for g in groups.values()) - len(out)


TABS = [  # tab name, criterion, question, loader
    ("C1 Unsourced statistics", "C1", "Statistic with no specific source nearby?", c1_citation_reason),
    ("C2 Wrong numbers", "C2", "Does a reliable source contradict the number?", c2_quantemp),
    ("C3 Loaded language", "C3", "Slanted / loaded wording where neutral would do?", c3_wnc_babe),
    ("C5 Vague attribution", "C5", "'Experts say' instead of a named source?", c5_szeged),
    ("C6 Causal claims", "C6", "Cause-and-effect stated without evidence?", c6_yu),
]
COLS = ["#", "Row ID", "Split", "Dataset", "Sentence", "Context / source", "Their label", "Expected",
        "Agent decision", "Correct?", "Error type", "Spot-check: your label", "Notes"]
WIDTHS = [5, 10, 7, 18, 70, 24, 26, 10, 12, 9, 14, 14, 26]


def _split(rows, rnd):
    """Dev/test split and spot-check rows; both halves of a pair stay together."""
    units = {}
    for r in rows:
        units.setdefault(r["pair"] or id(r), []).append(r)
    keys = list(units)
    rnd.shuffle(keys)
    n_dev = round(len(keys) * DEV_SHARE)
    ordered = []
    for i, k in enumerate(keys):
        for r in units[k]:
            r["split"] = "dev" if i < n_dev else "test"
            ordered.append(r)
    for r in rnd.sample(ordered, min(SPOT_CHECK, len(ordered))):
        r["spot"] = True
    return ordered


def _cell(ws, r, c, v, fill=None, bold=False, color=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font, cell.alignment, cell.border = Font(name=F, bold=bold, color=color), WRAP, THIN
    if fill:
        cell.fill = fill
    return cell


def build():
    rnd = random.Random(SEED)
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    samples, pools = {}, {}
    for name, cid, question, loader in TABS:
        rows, pool = loader(rnd)
        rows = _split(rows, rnd)
        pools[name] = pool
        ws = wb.create_sheet(name)
        _cell(ws, 1, 1, f"{cid}: {question}", bold=True)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLS))
        for c, (h, w) in enumerate(zip(COLS, WIDTHS), 1):
            _cell(ws, 2, c, h, HEAD, bold=True, color="FFFFFF")
            ws.column_dimensions[ws.cell(row=2, column=c).column_letter].width = w
        ws.freeze_panes = "A3"
        for i, r in enumerate(rows, 1):
            x = i + 2
            r["id"] = f"{cid}-{i:03d}"
            vals = [i, r["id"], r["split"], r["dataset"], r["sentence"], r["context"], r["their_label"], r["expected"],
                    None, f'=IF(I{x}="","",IF((I{x}="met")=(H{x}="met"),1,0))', None, None, None]
            for c, v in enumerate(vals, 1):
                _cell(ws, x, c, v, INPUT if (c == 12 and r.get("spot")) else None)
        samples[cid] = [{k: r[k] for k in ("id", "split", "dataset", "sentence", "context", "their_label",
                                           "expected", "pair")} for r in rows]

    # Summary: what each tab holds, then GiSuN's scores once the Agent decision column is filled
    summary.column_dimensions["A"].width = 26
    for col in "BCDEFGHIJKLM":
        summary.column_dimensions[col].width = 12
    intro = [
        "GiSuN on public datasets: one tab per criterion. Nothing has been run yet.",
        "Expected = the dataset's own (human) label mapped to our met / not_met. Sources and mappings: "
        "Google Doc 'GiSuN Public Datasets: Download Sources'.",
        "Split: dev rows are for tuning GiSuN, test rows are for the final numbers (run once). Rows not in this "
        "workbook stay in the raw files as the training pool.",
        "Yellow cells (30 per tab): one of you labels these by hand, without looking at Expected, to check the mapping.",
        "Agent decision and Error type are filled in after the GiSuN run; Correct? and this table update by themselves.",
    ]
    for i, t in enumerate(intro, 1):
        _cell(summary, i, 1, t, bold=(i == 1))
        summary.merge_cells(start_row=i, start_column=1, end_row=i, end_column=13)
    hdr = ["Tab", "Rows", "Dev", "Test", "Expected met", "Training pool left", "Agent decided", "TP", "FP", "FN",
           "Precision", "Recall", "F1", "Spot-checked", "Spot-check agrees"]
    h = len(intro) + 2
    for c, v in enumerate(hdr, 1):
        _cell(summary, h, c, v, HEAD, bold=True, color="FFFFFF")
    for i, (name, _, _, _) in enumerate(TABS):
        r = h + 1 + i
        q = f"'{name}'!"
        col = lambda L: f"{q}${L}$3:${L}$1000"
        vals = [name, f'=COUNTA({col("B")})', f'=COUNTIF({col("C")},"dev")', f'=COUNTIF({col("C")},"test")',
                f'=COUNTIF({col("H")},"met")', pools[name], f'=COUNTA({col("I")})',
                f'=COUNTIFS({col("H")},"met",{col("I")},"met")', f'=COUNTIFS({col("H")},"not_met",{col("I")},"met")',
                f'=COUNTIFS({col("H")},"met",{col("I")},"not_met")+COUNTIFS({col("H")},"met",{col("I")},"undetermined")',
                f'=IFERROR(H{r}/(H{r}+I{r}),"")', f'=IFERROR(H{r}/(H{r}+J{r}),"")',
                f'=IFERROR(2*K{r}*L{r}/(K{r}+L{r}),"")', f'=COUNTA({col("L")})',
                f'=IFERROR(SUMPRODUCT(({col("L")}<>"")*({col("L")}={col("H")}))/N{r},"")']
        for c, v in enumerate(vals, 1):
            cell = _cell(summary, r, c, v)
            if c in (11, 12, 13, 15):
                cell.number_format = "0.00"
    wb.save(OUT)
    json.dump(samples, open(os.path.join(X, "samples.json"), "w"), indent=1, ensure_ascii=False)
    return {name: len(samples[cid]) for name, cid, _, _ in TABS}, pools


if __name__ == "__main__":
    counts, pools = build()
    for k, v in counts.items():
        print(f"{k}: {v} rows (training pool left: {pools[k]})")
    print("wrote", OUT)
