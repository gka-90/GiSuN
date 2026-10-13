"""
Build an Excel worksheet that puts the rules_only and agent decisions side by side for every
(answer, sentence, criterion) and leaves a column for a human label.

    python -m evaluation.make_worksheet \
        --rules test_reports/chatgpt_rules_only.json \
        --agent test_reports/chatgpt_full_Qwen2.5-32B-Instruct-AWQ.json \
        --answers datasets/candidates_wildchat.json datasets/candidates_hc3.json \
        --out test_reports/compare_worksheet.xlsx

Rows where the two systems disagree come first. Fill the yellow "Human label" column
(met / not_met / unsure); the Summary sheet updates by formula.
"""

import argparse
import json

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

CRIT = {"C1": "Statistic with no specific source nearby?",
        "C2": "Does a reliable source contradict the number?",
        "C3": "Slanted / loaded wording where neutral would do?",
        "C4": "Sweeping absolute claim the evidence can't support?",
        "C5": "'Experts say' instead of a named source?",
        "C6": "Cause-and-effect stated without evidence?"}
F = "Arial"


def decisions(path):
    out = {}
    for r in json.load(open(path, encoding="utf-8"))["answers_detail"]:
        for c in r["claims"]:
            for k, v in c["criteria"].items():
                out[(r["id"], c["sentence"], k)] = (r.get("domain", ""), v)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--rules", required=True)
    ap.add_argument("--agent", required=True)
    ap.add_argument("--answers", nargs="*", default=[])
    ap.add_argument("--out", required=True)
    a = ap.parse_args()

    rules, agent = decisions(a.rules), decisions(a.agent)
    questions = {}
    for p in a.answers:
        for it in json.load(open(p, encoding="utf-8")):
            questions[it["id"]] = it.get("question", "")

    rows = []
    for key in set(rules) | set(agent):
        aid, sent, crit = key
        dom = (rules.get(key) or agent.get(key))[0]
        rv = rules.get(key, (None, "not run"))[1]
        av = agent.get(key, (None, "not run"))[1]
        rows.append((rv == av, aid, crit, dom, sent, rv, av))
    rows.sort(key=lambda r: (r[0], r[1], r[2]))  # disagreements first

    wb = Workbook()
    ws = wb.active
    ws.title = "Label"
    head = ["#", "Answer ID", "Domain", "Question", "Sentence", "Criterion", "Criterion question",
            "Rules decision", "Agent decision", "Agree?", "Human label", "Notes",
            "Rules correct?", "Agent correct?"]
    ws.append(head)
    for i, (agree, aid, crit, dom, sent, rv, av) in enumerate(rows, start=2):
        ws.append([i - 1, aid, dom, questions.get(aid, ""), sent, crit, CRIT.get(crit, ""), rv, av,
                   f'=IF(H{i}=I{i},"agree","DISAGREE")', None, None,
                   f'=IF(K{i}="","",IF(K{i}="unsure","",IF(H{i}=K{i},1,0)))',
                   f'=IF(K{i}="","",IF(K{i}="unsure","",IF(I{i}=K{i},1,0)))'])
    n = len(rows) + 1

    thin = Side(style="thin", color="D9D9D9")
    hfill = PatternFill("solid", fgColor="1F3864")
    for c in ws[1]:
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = hfill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for row in ws.iter_rows(min_row=2, max_row=n):
        for c in row:
            c.font = Font(name=F, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=thin)
        row[10].fill = yellow
        row[11].fill = yellow
    for col, w in zip("ABCDEFGHIJKLMN", [5, 10, 11, 30, 60, 7, 26, 11, 11, 10, 12, 30, 9, 9]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:N{n}"
    ws["K1"].comment = Comment("Type or pick: met (problem present), not_met, or unsure", "GiSuN")

    dv = DataValidation(type="list", formula1='"met,not_met,unsure"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"K2:K{n}")
    red = PatternFill("solid", fgColor="F8CBAD")
    ws.conditional_formatting.add(f"J2:J{n}", FormulaRule(formula=[f'J2="DISAGREE"'], fill=red))

    s = wb.create_sheet("Summary")
    s.append(["Criterion", "Question", "Pairs", "Disagreements", "Human-labeled",
              "Rules accuracy", "Agent accuracy", "Rules says met", "Agent says met"])
    L = f"Label!$F$2:$F${n}"
    for j, crit in enumerate(sorted(CRIT), start=2):
        s.append([crit, CRIT[crit],
                  f'=COUNTIF({L},A{j})',
                  f'=COUNTIFS({L},A{j},Label!$J$2:$J${n},"DISAGREE")',
                  f'=COUNTIFS({L},A{j},Label!$M$2:$M${n},"<>")-COUNTIFS({L},A{j},Label!$M$2:$M${n},"")',
                  f'=IFERROR(SUMIFS(Label!$M$2:$M${n},{L},A{j})/E{j},"")',
                  f'=IFERROR(SUMIFS(Label!$N$2:$N${n},{L},A{j})/E{j},"")',
                  f'=COUNTIFS({L},A{j},Label!$H$2:$H${n},"met")',
                  f'=COUNTIFS({L},A{j},Label!$I$2:$I${n},"met")'])
    t = len(CRIT) + 2
    s.append(["All", "", f"=SUM(C2:C{t-1})", f"=SUM(D2:D{t-1})", f"=SUM(E2:E{t-1})",
              f'=IFERROR(SUMPRODUCT(F2:F{t-1},E2:E{t-1})/E{t},"")',
              f'=IFERROR(SUMPRODUCT(G2:G{t-1},E2:E{t-1})/E{t},"")',
              f"=SUM(H2:H{t-1})", f"=SUM(I2:I{t-1})"])
    for c in s[1]:
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = hfill
        c.alignment = Alignment(wrap_text=True)
    for row in s.iter_rows(min_row=2, max_row=t):
        for c in row:
            c.font = Font(name=F, bold=(row[0].value == "All"))
        row[5].number_format = row[6].number_format = "0.0%"
    for col, w in zip("ABCDEFGHI", [10, 46, 8, 13, 13, 13, 13, 13, 13]):
        s.column_dimensions[col].width = w
    notes = [
        "How to use",
        "1. On the Label sheet, fill the yellow 'Human label' column: met = the problem IS present, not_met = it is not, unsure = skip.",
        "2. Start at the top: rows where rules and agent DISAGREE come first; they decide which system is better.",
        "3. Accuracy here = share of human-labeled rows where the system matched you. 'not run' never matches a label.",
        "4. Label independently from your partner, then compare; don't look at the other person's labels first.",
        f"Sources: {a.rules}; {a.agent}. Rows: one per (answer, sentence, criterion) that either system evaluated.",
    ]
    for k, line in enumerate(notes):
        cell = s.cell(t + 2 + k, 1, line)
        cell.font = Font(name=F, bold=(k == 0), italic=(k == len(notes) - 1), color="404040")
    wb.move_sheet("Summary", offset=-1)
    wb.save(a.out)
    dis = sum(1 for r in rows if not r[0])
    print(f"{len(rows)} rows ({dis} disagreements) -> {a.out}")


if __name__ == "__main__":
    main()
